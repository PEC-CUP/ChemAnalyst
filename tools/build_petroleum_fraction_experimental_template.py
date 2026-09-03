from __future__ import annotations

import os
import shutil
from datetime import datetime
from pathlib import Path

import openpyxl
import pandas as pd
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


TARGET = Path("data/experimental_db/templates/chemanalyst_experimental_sample_template.xlsx")
FALLBACK_TARGET = TARGET.with_name("chemanalyst_experimental_sample_template_petroleum_fraction_updated.xlsx")
ELEMENTAL = Path(os.getenv("PETRO_FRACTION_ELEMENTAL_FILE", "data/private/elemental_composition.xlsx"))
HYDROCARBON = Path(os.getenv("PETRO_FRACTION_HYDROCARBON_FILE", "data/private/hydrocarbon_types.xlsx"))
HYDROCARBON_TYPES = [
    "paraffins",
    "monocycloparaffins",
    "dicycloparaffins",
    "tricycloparaffins",
    "tetracycloparaffins",
    "pentacycloparaffins",
    "hexacycloparaffins",
    "alkylbenzenes",
    "benzocycloparaffins",
    "benzodicycloparaffins",
    "naphthalenes",
    "acenaphthenes",
    "fluorenes",
    "phenanthrenes",
    "phenanthrocycloalkanes",
    "pyrenes",
    "chrysenes",
    "perylenes",
    "dibenzanthracenes",
]


def _safe_sheet_name(name: str) -> str:
    return name[:31]


def _style_sheet(ws) -> None:
    header_fill = PatternFill("solid", fgColor="D9E8FF")
    header_font = Font(bold=True, color="0F2442")
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
    ws.freeze_panes = "A2"
    for column_cells in ws.columns:
        max_len = 0
        column = column_cells[0].column
        for cell in column_cells:
            value = "" if cell.value is None else str(cell.value)
            max_len = max(max_len, len(value))
        ws.column_dimensions[get_column_letter(column)].width = min(max(max_len + 2, 12), 42)


def _append_rows(ws, rows: list[list[object]]) -> None:
    for row in rows:
        ws.append(row)
    _style_sheet(ws)


def _property_meta(column: str) -> tuple[str, str, str, str, str]:
    token = column.strip()
    lower = token.lower()
    if token == "Density":
        return "density_20c", "g/cm3", "bulk_property", "20 deg C", "Density at 20 deg C"
    if token == "Kinematic viscosity":
        return "kinematic_viscosity_50c", "mm2/s", "bulk_property", "50 deg C", "Kinematic viscosity at 50 deg C"
    if token == "Solidification point":
        return "solidification_point", "deg C", "bulk_property", "", "Solidification point"
    if token == "Wax content":
        return "wax_content", "wt_percent", "bulk_property", "", "Wax content"
    if token in {"Saturates", "Aromatics", "Resins"}:
        return token.lower(), "wt_percent", "sara", "", "SARA fraction"
    if token in {"C wt%", "H wt%", "O wt%", "N wt%", "S wt%"}:
        return token.replace(" wt%", "").strip(), "wt_percent", "elemental", "elemental composition", token
    if token == "H/C":
        return "H_C", "atomic_ratio", "elemental_ratio", "calculated", "Hydrogen-to-carbon ratio"
    if lower in {"sat/aro", "sat/ar"}:
        return "SAT_AR", "ratio", "composition_ratio", "calculated", "Saturates-to-aromatics ratio"
    if lower == "cnn/bnn":
        return "", "", "", "", ""
    if token.startswith("BP_"):
        return token, "deg C", "boiling_point_distribution", "boiling point-yield distribution", f"Boiling point at {token.split('_', 1)[1]}% yield"
    return token, "", "bulk_property", "", token


def build_template() -> None:
    if not ELEMENTAL.exists():
        raise FileNotFoundError(ELEMENTAL)
    if not HYDROCARBON.exists():
        raise FileNotFoundError(HYDROCARBON)
    TARGET.parent.mkdir(parents=True, exist_ok=True)
    if TARGET.exists():
        backup = TARGET.with_name(f"{TARGET.stem}.backup_{datetime.now():%Y%m%d_%H%M%S}{TARGET.suffix}")
        shutil.copy2(TARGET, backup)

    elemental_df = pd.read_excel(ELEMENTAL)
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    ws = wb.create_sheet("sample_information")
    _append_rows(
        ws,
        [
            ["field", "value"],
            ["sample_id", "petroleum_fraction_db_001"],
            ["sample_name", "petroleum fraction 01"],
            ["sample_type", "petroleum fraction"],
            ["refinery", ""],
            ["project", "petroleum fraction historical experimental database"],
            ["raw_sample_id", "petroleum_fraction_01"],
            ["related_sample_id", ""],
            ["relation_type", ""],
            ["relation_description", ""],
        ],
    )

    ws = wb.create_sheet("bulk_properties")
    property_rows = [["property_name", "value", "unit", "method", "property_group", "source_column", "description"]]
    seen_properties: set[str] = set()
    for column in elemental_df.columns:
        if str(column).strip() == "Sample":
            continue
        name, unit, group, method, description = _property_meta(str(column))
        if not name:
            continue
        if name in seen_properties:
            continue
        seen_properties.add(name)
        property_rows.append([name, "", unit, method, group, str(column), description])
    _append_rows(ws, property_rows)

    for sheet_name in ("gc_fid", "gc_scd", "gc_ncd"):
        ws = wb.create_sheet(sheet_name)
        _append_rows(
            ws,
            [
                ["time", "intensity"],
                [0.000333333, ""],
            ],
        )

    ws = wb.create_sheet("ir_spectrum")
    _append_rows(
        ws,
        [
            ["wave_number_cm-1", "transmittance"],
            [3998.356867, ""],
        ],
    )

    ws = wb.create_sheet("hydrocarbon_types")
    hydro_rows = [["component", "content", "unit"]]
    for component in HYDROCARBON_TYPES:
        hydro_rows.append([component, "", "wt_percent"])
    _append_rows(ws, hydro_rows)

    ws = wb.create_sheet("sample_relations")
    _append_rows(
        ws,
        [
            ["related_sample_id", "relation_type", "description"],
            ["", "", ""],
        ],
    )

    try:
        wb.save(TARGET)
        print(TARGET)
    except PermissionError:
        wb.save(FALLBACK_TARGET)
        print(FALLBACK_TARGET)


if __name__ == "__main__":
    build_template()
