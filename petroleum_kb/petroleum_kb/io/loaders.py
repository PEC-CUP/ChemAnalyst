from __future__ import annotations

import json
from pathlib import Path

from ..schemas import LoadedDocument, SourceRegistration


SUPPORTED_SUFFIXES = {".txt", ".md", ".json", ".jsonl", ".pdf", ".docx"}


def load_documents(input_dir: Path, *, pdf_parser_backend: str = "pymupdf", grobid_url: str = "http://localhost:8070") -> list[LoadedDocument]:
    return load_documents_from_dirs([input_dir], pdf_parser_backend=pdf_parser_backend, grobid_url=grobid_url)


def load_documents_from_dirs(
    input_dirs: list[Path],
    *,
    pdf_parser_backend: str = "pymupdf",
    grobid_url: str = "http://localhost:8070",
    grobid_timeout_seconds: int = 300,
    strict_pdf_parser: bool = False,
    errors: list[dict] | None = None,
) -> list[LoadedDocument]:
    documents: list[LoadedDocument] = []
    paths: list[Path] = []
    for root in input_dirs:
        # Allow passing a single file path via --input_dir for convenience.
        if root.is_file():
            if root.suffix.lower() in SUPPORTED_SUFFIXES:
                paths.append(root)
            continue
        for path in sorted(root.rglob("*")):
            if not path.is_file():
                continue
            if path.suffix.lower() not in SUPPORTED_SUFFIXES:
                continue
            paths.append(path)

    total = len(paths)
    for idx, path in enumerate(paths, start=1):
        print(f"[loaders] ({idx}/{total}) loading: {path.name}")
        try:
            documents.extend(
                load_document(
                    path,
                    pdf_parser_backend=pdf_parser_backend,
                    grobid_url=grobid_url,
                    grobid_timeout_seconds=grobid_timeout_seconds,
                    strict_pdf_parser=strict_pdf_parser,
                )
            )
        except Exception as exc:
            # Never abort a full ingest because a single file fails to parse.
            msg = str(exc)
            print(f"[loaders] ERROR loading {path}: {msg}")
            if errors is not None:
                errors.append(
                    {
                        "path": str(path),
                        "file_name": path.name,
                        "suffix": path.suffix.lower(),
                        "error": msg,
                        "pdf_parser_backend": pdf_parser_backend,
                    }
                )
    return documents


def load_documents_from_sources(
    sources: list[SourceRegistration],
    *,
    pdf_parser_backend: str = "pymupdf",
    grobid_url: str = "http://localhost:8070",
    grobid_timeout_seconds: int = 300,
    strict_pdf_parser: bool = False,
    errors: list[dict] | None = None,
) -> list[LoadedDocument]:
    documents: list[LoadedDocument] = []
    for source in sources:
        if not source.enabled:
            continue
        root = Path(source.path)
        if not root.exists():
            print(f"[loaders] Skip missing source directory: {root}")
            continue
        for path in sorted(root.rglob("*")):
            if not path.is_file():
                continue
            if path.suffix.lower() not in SUPPORTED_SUFFIXES:
                continue
            try:
                documents.extend(
                    _load_document(
                        path,
                        source_name=source.name,
                        source_group_hint=source.source_group_hint,
                        task_group_hint=source.task_group_hint,
                        pdf_parser_backend=pdf_parser_backend,
                        grobid_url=grobid_url,
                        grobid_timeout_seconds=grobid_timeout_seconds,
                        strict_pdf_parser=strict_pdf_parser,
                    )
                )
            except Exception as exc:
                msg = str(exc)
                print(f"[loaders] ERROR loading {path}: {msg}")
                if errors is not None:
                    errors.append(
                        {
                            "path": str(path),
                            "file_name": path.name,
                            "suffix": path.suffix.lower(),
                            "error": msg,
                            "pdf_parser_backend": pdf_parser_backend,
                            "source_name": source.name,
                        }
                    )
    return documents


def load_document(
    path: Path,
    *,
    pdf_parser_backend: str = "pymupdf",
    grobid_url: str = "http://localhost:8070",
    grobid_timeout_seconds: int = 300,
    strict_pdf_parser: bool = False,
) -> list[LoadedDocument]:
    return _load_document(
        path,
        source_name=None,
        source_group_hint=None,
        task_group_hint=None,
        pdf_parser_backend=pdf_parser_backend,
        grobid_url=grobid_url,
        grobid_timeout_seconds=grobid_timeout_seconds,
        strict_pdf_parser=strict_pdf_parser,
    )


def _load_document(
    path: Path,
    *,
    source_name: str | None,
    source_group_hint: str | None,
    task_group_hint: str | None,
    pdf_parser_backend: str = "pymupdf",
    grobid_url: str = "http://localhost:8070",
    grobid_timeout_seconds: int = 300,
    strict_pdf_parser: bool = False,
) -> list[LoadedDocument]:
    suffix = path.suffix.lower()
    if suffix in {".txt", ".md"}:
        return [_load_plain_text(path, source_name, source_group_hint, task_group_hint)]
    if suffix == ".json":
        return _load_json(path, source_name, source_group_hint, task_group_hint)
    if suffix == ".jsonl":
        return _load_jsonl(path, source_name, source_group_hint, task_group_hint)
    if suffix == ".pdf":
        return _load_pdf(
            path,
            source_name,
            source_group_hint,
            task_group_hint,
            pdf_parser_backend,
            grobid_url,
            grobid_timeout_seconds,
            strict_pdf_parser,
        )
    if suffix == ".docx":
        return _load_docx(path, source_name, source_group_hint, task_group_hint)
    return []


def _build_loaded_document(
    *,
    path: Path,
    text: str,
    source_name: str | None,
    source_group_hint: str | None,
    task_group_hint: str | None,
    page: int | None = None,
    section: str | None = None,
    title: str | None = None,
    authors: list[str] | None = None,
    year: int | None = None,
    journal: str | None = None,
    doi: str | None = None,
    abstract: str | None = None,
    document_type: str | None = None,
) -> LoadedDocument:
    return LoadedDocument(
        text=text,
        source=str(path),
        file_name=path.name,
        page=page,
        section=section,
        source_name=source_name,
        source_group_hint=source_group_hint,
        task_group_hint=task_group_hint,
        title=title,
        authors=authors or [],
        year=year,
        journal=journal,
        doi=doi,
        abstract=abstract,
        document_type=document_type,
    )


def _load_plain_text(
    path: Path,
    source_name: str | None,
    source_group_hint: str | None,
    task_group_hint: str | None,
) -> LoadedDocument:
    return _build_loaded_document(
        path=path,
        text=path.read_text(encoding="utf-8", errors="ignore"),
        source_name=source_name,
        source_group_hint=source_group_hint,
        task_group_hint=task_group_hint,
    )


def _load_json(
    path: Path,
    source_name: str | None,
    source_group_hint: str | None,
    task_group_hint: str | None,
) -> list[LoadedDocument]:
    payload = json.loads(path.read_text(encoding="utf-8", errors="ignore"))
    if isinstance(payload, dict):
        text = payload.get("text") or json.dumps(payload, ensure_ascii=False)
        return [
            _build_loaded_document(
                path=path,
                text=text,
                source_name=source_name,
                source_group_hint=source_group_hint,
                task_group_hint=task_group_hint,
            )
        ]
    if isinstance(payload, list):
        docs: list[LoadedDocument] = []
        for idx, item in enumerate(payload):
            if isinstance(item, dict):
                text = item.get("text") or json.dumps(item, ensure_ascii=False)
                docs.append(
                    _build_loaded_document(
                        path=path,
                        text=text,
                        source_name=source_name,
                        source_group_hint=source_group_hint,
                        task_group_hint=task_group_hint,
                        section=f"item_{idx}",
                    )
                )
        return docs
    return [
        _build_loaded_document(
            path=path,
            text=str(payload),
            source_name=source_name,
            source_group_hint=source_group_hint,
            task_group_hint=task_group_hint,
        )
    ]


def _load_jsonl(
    path: Path,
    source_name: str | None,
    source_group_hint: str | None,
    task_group_hint: str | None,
) -> list[LoadedDocument]:
    docs: list[LoadedDocument] = []
    with path.open("r", encoding="utf-8", errors="ignore") as fh:
        for idx, line in enumerate(fh):
            line = line.strip()
            if not line:
                continue
            try:
                item = json.loads(line)
            except json.JSONDecodeError:
                item = {"text": line}
            text = item.get("text") if isinstance(item, dict) else str(item)
            docs.append(
                _build_loaded_document(
                    path=path,
                    text=text,
                    source_name=source_name,
                    source_group_hint=source_group_hint,
                    task_group_hint=task_group_hint,
                    section=f"line_{idx}",
                )
            )
    return docs


def _load_pdf(
    path: Path,
    source_name: str | None,
    source_group_hint: str | None,
    task_group_hint: str | None,
    pdf_parser_backend: str,
    grobid_url: str,
    grobid_timeout_seconds: int,
    strict_pdf_parser: bool,
) -> list[LoadedDocument]:
    backend = (pdf_parser_backend or "").strip().lower()
    if backend == "grobid":
        # Strict mode: do not fall back to PyMuPDF when GROBID is selected.
        # If GROBID fails or returns empty content, we treat it as an ingest failure
        # and let the caller record the error and continue with other files.
        from .grobid_loader import load_pdf_with_grobid

        docs = load_pdf_with_grobid(
            path,
            source_name=source_name,
            source_group_hint=source_group_hint,
            task_group_hint=task_group_hint,
            grobid_url=grobid_url,
            timeout_seconds=grobid_timeout_seconds,
        )
        if not docs:
            raise RuntimeError("GROBID returned no extracted text")
        return docs

    if strict_pdf_parser:
        raise RuntimeError(f"strict_pdf_parser enabled: refusing pdf_parser_backend={backend!r} for {path.name}")

    try:
        import fitz  # type: ignore
    except ImportError:
        print(f"[loaders] Skip PDF because PyMuPDF is unavailable: {path}")
        return []

    docs: list[LoadedDocument] = []
    with fitz.open(path) as pdf:
        for page_index, page in enumerate(pdf, start=1):
            text = page.get_text("text").strip()
            if not text:
                continue
            docs.append(
                _build_loaded_document(
                    path=path,
                    text=text,
                    source_name=source_name,
                    source_group_hint=source_group_hint,
                    task_group_hint=task_group_hint,
                    page=page_index,
                    section=f"page_{page_index}",
                    document_type="pdf",
                )
            )
    return docs


def _load_docx(
    path: Path,
    source_name: str | None,
    source_group_hint: str | None,
    task_group_hint: str | None,
) -> list[LoadedDocument]:
    try:
        from docx import Document  # type: ignore
    except ImportError:
        print(f"[loaders] Skip DOCX because python-docx is unavailable: {path}")
        return []

    doc = Document(path)
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    text = "\n".join(paragraphs)
    if not text:
        return []
    return [
        _build_loaded_document(
            path=path,
            text=text,
            source_name=source_name,
            source_group_hint=source_group_hint,
            task_group_hint=task_group_hint,
        )
    ]
