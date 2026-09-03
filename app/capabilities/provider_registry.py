from __future__ import annotations

import importlib
import json
import re
import shutil
from pathlib import Path
from typing import Any, Callable

from app.agents.schemas import CapabilitySchema
from app.agents.task_specific_evaluation import (
    evaluate_tool_provider_result,
    summarize_task_specific_results,
    task_specific_cases_from_payload,
)
from app.tools.adapters.generic_external_python_adapter import build_adapter as build_generic_python_adapter
from app.tools.registry import CapabilityRegistry


PROVIDER_ROOT = Path(__file__).resolve().parent / "providers"


def save_provider_package(package: dict[str, Any], *, provider_id: str | None = None, root: Path | None = None) -> dict[str, Any]:
    root = root or PROVIDER_ROOT
    capability = package.get("capability") if isinstance(package, dict) else {}
    manifest = package.get("mcp_manifest") if isinstance(package, dict) else {}
    capability_id = str((capability or {}).get("capability_id") or (manifest or {}).get("capability_id") or provider_id or "unnamed")
    provider = _slugify(provider_id or capability_id.replace("tool.", ""))
    provider_dir = root / provider
    provider_dir.mkdir(parents=True, exist_ok=True)

    manifest = dict(manifest or {})
    manifest.setdefault("schema_version", "chemanalyst-mcp-tool.v1")
    manifest.setdefault("capability_id", capability_id)
    manifest.setdefault("review", {"reviewed": False, "enabled": False})
    manifest.setdefault("provider_id", provider)
    _ensure_input_templates(provider_dir, capability, manifest)
    package = dict(package or {})
    package["capability"] = capability
    package["mcp_manifest"] = manifest

    files = {
        "package.json": package,
        "capability.json": capability,
        "manifest.json": manifest,
        "evaluation_cases.json": package.get("evaluation_cases", []),
        "registration_plan.json": package.get("registration_plan", {}),
    }
    for name, payload in files.items():
        (provider_dir / name).write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    return {
        "provider_id": provider,
        "provider_dir": str(provider_dir),
        "capability_id": capability_id,
        "saved_files": sorted(files.keys()),
        "reviewed": False,
        "enabled": False,
    }


def list_provider_packages(*, root: Path | None = None) -> list[dict[str, Any]]:
    root = root or PROVIDER_ROOT
    if not root.exists():
        return []
    providers = []
    for provider_dir in sorted(path for path in root.iterdir() if path.is_dir()):
        manifest = _read_json(provider_dir / "manifest.json")
        capability = _read_json(provider_dir / "capability.json")
        providers.append(
            {
                "provider_id": provider_dir.name,
                "capability_id": capability.get("capability_id") or manifest.get("capability_id"),
                "description": capability.get("description") or manifest.get("description"),
                "review": manifest.get("review") or {},
                "workflow_relations": manifest.get("workflow_relations") or {},
            }
        )
    return providers


def provider_runtime_spec(provider_id: str, *, root: Path | None = None) -> dict[str, Any]:
    provider = _slugify(provider_id)
    provider_dir = (root or PROVIDER_ROOT) / provider
    capability = _read_json(provider_dir / "capability.json")
    manifest = _read_json(provider_dir / "manifest.json")
    if not capability and not manifest:
        return {"status": "error", "message": "provider_not_found", "provider_id": provider}

    _materialize_input_template_paths(capability, provider_dir)
    input_schema = capability.get("input_schema") if isinstance(capability.get("input_schema"), dict) else {}
    requirements = input_schema.get("x_chemanalyst_input_requirements")
    if not isinstance(requirements, dict):
        requirements = {}
    templates = requirements.get("input_templates") if isinstance(requirements.get("input_templates"), list) else []
    template_input_file = None
    for template in templates:
        if isinstance(template, dict) and template.get("path"):
            template_input_file = str(template["path"])
            break
    review = manifest.get("review") if isinstance(manifest.get("review"), dict) else {}
    capability_id = str(capability.get("capability_id") or manifest.get("capability_id") or "")
    return {
        "status": "success",
        "message": "provider_runtime_spec",
        "provider_id": provider,
        "provider_dir": str(provider_dir),
        "capability_id": capability_id,
        "tool_name": capability_id,
        "review": review,
        "callable": bool(review.get("reviewed") and review.get("enabled")),
        "description": capability.get("description") or manifest.get("description"),
        "input_requirements": requirements,
        "template_input_file": template_input_file,
        "capability": capability,
        "manifest": manifest,
    }


def evaluate_provider_package(
    provider_id: str,
    *,
    execute: bool = False,
    sample_payload: dict[str, Any] | None = None,
    root: Path | None = None,
) -> dict[str, Any]:
    provider_dir = (root or PROVIDER_ROOT) / _slugify(provider_id)
    capability = _read_json(provider_dir / "capability.json")
    manifest = _read_json(provider_dir / "manifest.json")
    cases = _read_json(provider_dir / "evaluation_cases.json")
    checks = {
        "provider_dir_exists": provider_dir.exists(),
        "capability_schema_present": bool(capability.get("capability_id")),
        "mcp_manifest_present": bool(manifest.get("capability_id")),
        "evaluation_cases_present": isinstance(cases, list) and bool(cases),
        "evidence_contract_present": bool(manifest.get("evidence_contract")),
        "workflow_relations_present": "workflow_relations" in manifest,
    }
    execution_result: dict[str, Any] | None = None
    if execute:
        adapter = load_reviewed_adapter(provider_id, root=root)
        if adapter is None:
            checks["reviewed_adapter_available"] = False
            execution_result = {"status": "skipped", "message": "reviewed_adapter_not_available"}
        else:
            checks["reviewed_adapter_available"] = True
            try:
                payload = sample_payload or {}
                execution_result = adapter(payload)
                checks["execution_success"] = isinstance(execution_result, dict) and execution_result.get("status") == "success"
                checks["structured_evidence_present"] = _has_required_outputs(
                    execution_result if isinstance(execution_result, dict) else {},
                    manifest.get("evidence_contract", {}).get("required_outputs", []),
                )
            except Exception as exc:
                checks["execution_success"] = False
                execution_result = {"status": "error", "message": str(exc)}

    evaluation = {
        "provider_id": _slugify(provider_id),
        "checks": checks,
        "passed": all(checks.values()),
        "execution_result": execution_result,
        "capability": capability,
        "manifest": manifest,
        "evaluation_case_count": len(cases) if isinstance(cases, list) else 0,
    }
    task_specific_cases = [case for case in task_specific_cases_from_payload(cases) if case.evaluation_target == "tool"]
    task_specific_results = [evaluate_tool_provider_result(case, evaluation) for case in task_specific_cases]
    evaluation["task_specific_evaluation"] = {
        "stage": "offline_development",
        "case_source": "provider_evaluation_cases",
        "cases": [case.to_dict() for case in task_specific_cases],
        "results": [result.to_dict() for result in task_specific_results],
        "summary": summarize_task_specific_results(task_specific_results),
    }
    return evaluation


def register_reviewed_provider(
    registry: CapabilityRegistry,
    provider_id: str,
    *,
    root: Path | None = None,
) -> dict[str, Any]:
    provider_dir = (root or PROVIDER_ROOT) / _slugify(provider_id)
    capability_payload = _read_json(provider_dir / "capability.json")
    _materialize_input_template_paths(capability_payload, provider_dir)
    adapter = load_reviewed_adapter(provider_id, root=root)
    if adapter is None:
        return {"status": "error", "message": "reviewed_adapter_not_available", "provider_id": _slugify(provider_id)}
    capability = CapabilitySchema(**capability_payload)
    registry.register(capability, adapter)
    return {"status": "success", "message": "provider_registered", "provider_id": _slugify(provider_id), "capability_id": capability.capability_id}


def revoke_provider(
    registry: CapabilityRegistry,
    provider_id: str,
    *,
    reason: str | None = None,
    root: Path | None = None,
) -> dict[str, Any]:
    provider = _slugify(provider_id)
    provider_dir = (root or PROVIDER_ROOT) / provider
    capability_payload = _read_json(provider_dir / "capability.json")
    manifest_path = provider_dir / "manifest.json"
    manifest = _read_json(manifest_path)
    if not capability_payload and not manifest:
        return {"status": "error", "message": "provider_not_found", "provider_id": provider}
    capability_id = str(capability_payload.get("capability_id") or manifest.get("capability_id") or "")
    if manifest:
        manifest["review"] = {
            "reviewed": bool((manifest.get("review") or {}).get("reviewed")),
            "enabled": False,
            "notes": reason or "Provider revoked from Tool Onboarding UI.",
        }
        manifest["revoked"] = True
        manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    unregistered = registry.unregister(capability_id) if capability_id and hasattr(registry, "unregister") else False
    return {
        "status": "success",
        "message": "provider_revoked",
        "provider_id": provider,
        "capability_id": capability_id,
        "runtime_unregistered": unregistered,
        "provider_files_kept": True,
    }


def update_provider_review(
    provider_id: str,
    *,
    reviewed: bool,
    enabled: bool,
    notes: str | None = None,
    root: Path | None = None,
) -> dict[str, Any]:
    provider_dir = (root or PROVIDER_ROOT) / _slugify(provider_id)
    manifest_path = provider_dir / "manifest.json"
    manifest = _read_json(manifest_path)
    if not manifest:
        return {"status": "error", "message": "provider_manifest_not_found", "provider_id": _slugify(provider_id)}
    manifest["review"] = {"reviewed": reviewed, "enabled": enabled, "notes": notes or ""}
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    return {
        "status": "success",
        "message": "provider_review_updated",
        "provider_id": _slugify(provider_id),
        "review": manifest["review"],
    }


def load_reviewed_adapter(provider_id: str, *, root: Path | None = None) -> Callable[[dict], dict] | None:
    manifest = _read_json((root or PROVIDER_ROOT) / _slugify(provider_id) / "manifest.json")
    review = manifest.get("review") if isinstance(manifest.get("review"), dict) else {}
    if not (review.get("reviewed") and review.get("enabled")):
        return None
    invocation = manifest.get("invocation") if isinstance(manifest.get("invocation"), dict) else {}
    if invocation.get("type") == "python_script" and invocation.get("adapter_template") == "generic_python_function":
        return build_generic_python_adapter(manifest)
    if invocation.get("type") != "python_callable":
        return None
    entrypoint = str(invocation.get("entrypoint") or "")
    if not entrypoint.startswith("app.tools."):
        return None
    module_name, _, attr = entrypoint.rpartition(".")
    if not module_name or not attr:
        return None
    module = importlib.import_module(module_name)
    func = getattr(module, attr, None)
    return func if callable(func) else None


def _has_required_outputs(payload: dict[str, Any], required_outputs: list[str]) -> bool:
    if not required_outputs:
        return True
    return all(output in payload or output in (payload.get("result") or {}) for output in required_outputs)


def _read_json(path: Path) -> Any:
    if not path.exists():
        return {}
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except Exception:
        return {}


def _ensure_input_templates(provider_dir: Path, capability: dict[str, Any], manifest: dict[str, Any]) -> None:
    input_schema = capability.get("input_schema") if isinstance(capability.get("input_schema"), dict) else {}
    requirements = input_schema.get("x_chemanalyst_input_requirements")
    if not isinstance(requirements, dict):
        return
    if requirements.get("input_templates"):
        _copy_declared_input_templates(provider_dir, requirements)
        return
    required_sheets = requirements.get("required_sheets")
    if not isinstance(required_sheets, dict) or not required_sheets:
        return
    template_dir = provider_dir / "templates"
    template_dir.mkdir(parents=True, exist_ok=True)
    template_path = template_dir / "input_template.xlsx"
    _write_excel_input_template(template_path, required_sheets, requirements)
    requirements["input_templates"] = [
        {
            "label": "Example input workbook",
            "field": "input_file",
            "path": "templates/input_template.xlsx",
            "description": "Excel workbook generated from the onboarded tool input schema with example numeric values.",
        }
    ]
    manifest["input_schema"] = input_schema


def _copy_declared_input_templates(provider_dir: Path, requirements: dict[str, Any]) -> None:
    templates = requirements.get("input_templates")
    if not isinstance(templates, list):
        return
    for template in templates:
        if not isinstance(template, dict):
            continue
        source_path = str(template.get("source_path") or "").strip()
        if not source_path:
            continue
        source = Path(source_path)
        if not source.exists() or not source.is_file():
            continue
        raw_dest = str(template.get("path") or "").strip() or f"templates/{source.name}"
        dest = Path(raw_dest)
        if not dest.is_absolute():
            dest = provider_dir / dest
        dest.parent.mkdir(parents=True, exist_ok=True)
        shutil.copyfile(source, dest)
        try:
            template["path"] = str(dest.relative_to(provider_dir))
        except ValueError:
            template["path"] = str(dest)


def _write_excel_input_template(template_path: Path, required_sheets: dict[str, Any], requirements: dict[str, Any]) -> None:
    from openpyxl import Workbook

    wb = Workbook()
    first_sheet = True
    for sheet_name, columns in required_sheets.items():
        ws = wb.active if first_sheet else wb.create_sheet()
        first_sheet = False
        ws.title = str(sheet_name)[:31] or "input"
        if isinstance(columns, list):
            normalized_columns = [str(column) for column in columns]
            ws.append(normalized_columns)
            for row in _example_rows_for_input_sheet(str(sheet_name), normalized_columns):
                ws.append(row)
    notes = requirements.get("notes")
    if isinstance(notes, list) and notes:
        ws = wb.create_sheet("README")
        ws.append(["Input requirement notes"])
        for note in notes:
            ws.append([str(note)])
    wb.save(template_path)


def _example_rows_for_input_sheet(sheet_name: str, columns: list[str]) -> list[list[Any]]:
    normalized_sheet = sheet_name.strip().lower()
    normalized_columns = [column.strip().lower() for column in columns]
    if normalized_sheet == "fractions":
        row = []
        for column in normalized_columns:
            if "saturate" in column:
                row.append(62.5)
            elif "aromatic" in column:
                row.append(37.5)
            else:
                row.append(1.0)
        return [row]
    if "m/z" in normalized_columns and "intensity" in normalized_columns:
        mz_idx = normalized_columns.index("m/z")
        intensity_idx = normalized_columns.index("intensity")
        seeds = [(178.0, 12850.0), (192.0, 21400.0), (206.0, 17320.0), (220.0, 9410.0)]
        rows = []
        for mz, intensity in seeds:
            row = [None for _ in columns]
            row[mz_idx] = mz
            row[intensity_idx] = intensity
            rows.append(row)
        return rows
    return [[float(index + 1) for index, _column in enumerate(columns)]]


def _materialize_input_template_paths(capability_payload: dict[str, Any], provider_dir: Path) -> None:
    input_schema = capability_payload.get("input_schema") if isinstance(capability_payload.get("input_schema"), dict) else {}
    requirements = input_schema.get("x_chemanalyst_input_requirements")
    if not isinstance(requirements, dict):
        return
    templates = requirements.get("input_templates")
    if not isinstance(templates, list):
        return
    for template in templates:
        if not isinstance(template, dict):
            continue
        raw_path = str(template.get("path") or "").strip()
        if not raw_path:
            continue
        path = Path(raw_path)
        if not path.is_absolute():
            path = provider_dir / path
        template["path"] = str(path.resolve())


def _slugify(value: str) -> str:
    cleaned = re.sub(r"[^a-zA-Z0-9]+", "_", str(value).strip().lower()).strip("_")
    return cleaned or "unnamed_provider"
