from __future__ import annotations

import os
import csv
import json
import random
import re
import shutil
from datetime import datetime
from pathlib import Path
from typing import Any

import openpyxl
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


SOURCE_ROOT = Path(os.getenv("PETRO_FRACTION_SOURCE_ROOT", "data/private/petroleum_fraction_source"))


def _legacy_source_workbook(channel: str) -> Path:
    # Keep compatibility with the original raw-data filenames without exposing that
    # legacy sample-type abbreviation as a ChemAnalyst identifier.
    return SOURCE_ROOT / "GC" / ("V" + f"GO_{channel}.xlsx")


GC_FID = _legacy_source_workbook("GC")
GC_SCD = _legacy_source_workbook("SCD")
GC_NCD = _legacy_source_workbook("NCD")
IR = SOURCE_ROOT / "common data" / "IR.xlsx"
HYDROCARBON = SOURCE_ROOT / "common data" / "hydrocarbon_types.xlsx"
ELEMENTAL = SOURCE_ROOT / "elemental composition.xlsx"
QUANTIFIED_DIR = SOURCE_ROOT / "fitted & quantified composition" / "quantified" / "experimental_composition"

OUT_ROOT = Path(os.getenv("PETRO_FRACTION_WORKBOOK_OUT", "data/experimental_db/import_workbooks"))
TEMPLATE = Path(os.getenv("PETRO_FRACTION_TEMPLATE", "data/experimental_db/templates/chemanalyst_experimental_sample_template.xlsx"))
RANDOM_SEED = 20260711
TEST_FRACTION = 0.20

GC_FID_FALLBACKS = {
    "petroleum_fraction_22": "petroleum_fraction_21",
}

HYDROCARBON_TYPES = [
    "paraffins",
    "monocycloparaffins",
    "dicycloparaffins",
    "tricycloparaffins",
    "tetracycloparaffins",
    "pentacycloparaffins",
    "hexacycloparaffins",
    "alkylbenzenes",
    "benzocycloparaffins",
    "benzodicycloparaffins",
    "naphthalenes",
    "acenaphthenes",
    "fluorenes",
    "phenanthrenes",
    "phenanthrocycloalkanes",
    "pyrenes",
    "chrysenes",
    "perylenes",
    "dibenzanthracenes",
]

SATURATED_DBE_TO_TYPE = {
    0: "paraffins",
    1: "monocycloparaffins",
    2: "dicycloparaffins",
    3: "tricycloparaffins",
    4: "tetracycloparaffins",
    5: "pentacycloparaffins",
    6: "hexacycloparaffins",
}

AROMATIC_DBE_TO_TYPE = {
    4: "alkylbenzenes",
    5: "benzocycloparaffins",
    6: "benzodicycloparaffins",
    7: "naphthalenes",
    8: "acenaphthenes",
    9: "fluorenes",
    10: "phenanthrenes",
    11: "phenanthrocycloalkanes",
    12: "pyrenes",
    13: "chrysenes",
    14: "perylenes",
    15: "dibenzanthracenes",
}


def canonical_raw_id(value: Any) -> str:
    text = str(value or "").strip()
    match = re.fullmatch(r"(petroleum_fraction_62)[_-]([12])", text, flags=re.IGNORECASE)
    if match:
        return f"petroleum_fraction_62-{match.group(2)}"
    return text


def sample_sort_key(sample_id: str) -> tuple[int, int, str]:
    text = canonical_raw_id(sample_id)
    match = re.fullmatch(r"petroleum_fraction_(\d+)(top:-([12]))top", text, flags=re.IGNORECASE)
    if not match:
        return (9999, 0, text)
    return (int(match.group(1)), int(match.group(2) or 0), text)


def db_id_for(index: int) -> str:
    return f"petroleum_fraction_db_{index:03d}"


def sample_name_for(raw_id: str) -> str:
    return "petroleum fraction " + canonical_raw_id(raw_id).replace("petroleum_fraction_", "").replace("-", "-")


def ensure_output_root() -> None:
    OUT_ROOT.mkdir(parents=True, exist_ok=True)
    for child in ("train", "test"):
        path = OUT_ROOT / child
        if path.exists():
            backup = OUT_ROOT / f"{child}.backup_{datetime.now():%Y%m%d_%H%M%S}"
            shutil.move(str(path), str(backup))
        path.mkdir(parents=True, exist_ok=True)


def style_sheet(ws) -> None:
    header_fill = PatternFill("solid", fgColor="D9E8FF")
    header_font = Font(bold=True, color="0F2442")
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
    ws.freeze_panes = "A2"
    if ws.max_row > 5000:
        for column in range(1, ws.max_column + 1):
            ws.column_dimensions[get_column_letter(column)].width = 16
        return
    for column_cells in ws.columns:
        max_len = 0
        col = column_cells[0].column
        for cell in column_cells[:200]:
            value = "" if cell.value is None else str(cell.value)
            max_len = max(max_len, len(value))
        ws.column_dimensions[get_column_letter(col)].width = min(max(max_len + 2, 12), 36)


def append_rows(ws, rows: list[list[Any]]) -> None:
    for row in rows:
        ws.append(row)
    style_sheet(ws)


def property_meta(source_column: str) -> tuple[str, str, str, str, str] | None:
    token = str(source_column).strip()
    lower = token.lower()
    if token == "Sample":
        return None
    if token == "Density":
        return ("density_20c", "g/cm3", "20 deg C", "bulk_property", "Density at 20 deg C")
    if token == "Kinematic viscosity":
        return ("kinematic_viscosity_50c", "mm2/s", "50 deg C", "bulk_property", "Kinematic viscosity at 50 deg C")
    if token == "Solidification point":
        return ("solidification_point", "deg C", "", "bulk_property", "Solidification point")
    if token == "Wax content":
        return ("wax_content", "wt_percent", "", "bulk_property", "Wax content")
    if token in {"Saturates", "Aromatics", "Resins"}:
        return (token.lower(), "wt_percent", "", "sara", "SARA fraction")
    if token in {"C wt%", "H wt%", "O wt%", "N wt%", "S wt%"}:
        return (token.replace(" wt%", "").strip(), "wt_percent", "elemental composition", "elemental", token)
    if token == "H/C":
        return ("H_C", "atomic_ratio", "calculated", "elemental_ratio", "Hydrogen-to-carbon ratio")
    if lower in {"sat/aro", "sat/ar"}:
        return ("SAT_AR", "ratio", "calculated", "composition_ratio", "Saturates-to-aromatics ratio")
    if lower == "cnn/bnn":
        return None
    if token.startswith("BP_"):
        return (token, "deg C", "boiling point-yield distribution", "boiling_point_distribution", f"Boiling point at {token.split('_', 1)[1]}% yield")
    return (token, "", "", "bulk_property", token)


def numeric_or_blank(value: Any) -> Any:
    if pd.isna(value):
        return ""
    return value


def load_wide_table(path: Path) -> dict[str, pd.Series]:
    df = pd.read_excel(path)
    first_col = df.columns[0]
    out: dict[str, pd.Series] = {}
    for col in df.columns[1:]:
        out[canonical_raw_id(col)] = df[col]
    out["__x__"] = df[first_col]
    return out


def load_hydrocarbon_table() -> tuple[pd.Series, dict[str, pd.Series]]:
    df = pd.read_excel(HYDROCARBON)
    components = pd.Series(HYDROCARBON_TYPES)
    labels = df.iloc[:, 0].astype(str).str.strip()
    filtered = df.loc[labels.isin(HYDROCARBON_TYPES)].copy()
    filtered.insert(0, "__component__", labels[labels.isin(HYDROCARBON_TYPES)].values)
    filtered = filtered.set_index("__component__").reindex(HYDROCARBON_TYPES)
    values: dict[str, pd.Series] = {}
    for col in df.columns[1:]:
        values[canonical_raw_id(col)] = filtered[col].reset_index(drop=True)
    return components, values


def quantified_path_for(raw_id: str) -> Path:
    return QUANTIFIED_DIR / f"{raw_id.replace('-', '_')}_experimental composition_quantify.xlsx"


def hydrocarbon_from_quantified(raw_id: str) -> pd.Series | None:
    path = quantified_path_for(raw_id)
    if not path.exists():
        return None
    df = pd.read_excel(path, sheet_name="Quantified")
    compound_type = df.get("compound type")
    dbe = pd.to_numeric(df.get("DBE"), errors="coerce")
    mass_fraction = pd.to_numeric(df.get("mass_fraction"), errors="coerce")
    if compound_type is None:
        return None
    compound_type = compound_type.astype(str).str.strip()
    values = {name: 0.0 for name in HYDROCARBON_TYPES}
    for dbe_value, component in SATURATED_DBE_TO_TYPE.items():
        mask = compound_type.eq("Saturated hydrocarbons") & dbe.eq(dbe_value)
        values[component] = float(mass_fraction[mask].sum())
    for dbe_value, component in AROMATIC_DBE_TO_TYPE.items():
        mask = compound_type.eq("Aromatic hydrocarbons") & dbe.eq(dbe_value)
        values[component] = float(mass_fraction[mask].sum())
    return pd.Series([values[name] for name in HYDROCARBON_TYPES])


def chromatogram_source_id(sheet_name: str, raw_id: str) -> str:
    if sheet_name == "gc_fid":
        return GC_FID_FALLBACKS.get(raw_id, raw_id)
    return raw_id


def has_chromatogram(sheet_names: set[str], source_raw_id: str) -> bool:
    return source_raw_id in sheet_names or source_raw_id.replace("-", "_") in sheet_names


def chromatogram_rows(path: Path, raw_id: str) -> list[list[Any]]:
    xlsx = pd.ExcelFile(path)
    sheet = raw_id if raw_id in xlsx.sheet_names else raw_id.replace("-", "_")
    if sheet not in xlsx.sheet_names:
        return [["time", "intensity"]]
    frame = pd.read_excel(path, sheet_name=sheet, header=3, usecols=[0, 1])
    frame = frame.dropna(how="all")
    frame.columns = ["time", "intensity"]
    rows = [["time", "intensity"]]
    for time_value, intensity in frame.itertuples(index=False):
        if pd.isna(time_value) or pd.isna(intensity):
            continue
        rows.append([time_value, intensity])
    return rows


def write_sample_workbook(
    *,
    path: Path,
    sample_index: int,
    raw_id: str,
    elemental_row: pd.Series,
    ir_data: dict[str, pd.Series],
    hydro_components: pd.Series,
    hydro_values: dict[str, pd.Series],
    chromatogram_sheet_sets: dict[str, set[str]],
) -> dict[str, Any]:
    db_id = db_id_for(sample_index)
    wb = Workbook()
    wb.remove(wb.active)

    sample_info = [
        ["field", "value"],
        ["sample_id", db_id],
        ["sample_name", sample_name_for(raw_id)],
        ["sample_type", "petroleum fraction"],
        ["refinery", ""],
        ["project", "petroleum fraction historical experimental database"],
        ["raw_sample_id", raw_id],
        ["related_sample_id", ""],
        ["relation_type", ""],
        ["relation_description", ""],
    ]
    append_rows(wb.create_sheet("sample_information"), sample_info)

    seen_properties: set[str] = set()
    property_rows = [["property_name", "value", "unit", "method", "property_group", "source_column", "description"]]
    for col, value in elemental_row.items():
        meta = property_meta(str(col))
        if meta is None:
            continue
        name, unit, method, group, description = meta
        if name in seen_properties:
            continue
        seen_properties.add(name)
        property_rows.append([name, numeric_or_blank(value), unit, method, group, str(col), description])
    append_rows(wb.create_sheet("bulk_properties"), property_rows)

    for sheet_name, source_path in (
        ("gc_fid", GC_FID),
        ("gc_scd", GC_SCD),
        ("gc_ncd", GC_NCD),
    ):
        source_raw_id = chromatogram_source_id(sheet_name, raw_id)
        append_rows(wb.create_sheet(sheet_name), chromatogram_rows(source_path, source_raw_id))

    ir_rows = [["wave_number_cm-1", "transmittance"]]
    if raw_id in ir_data:
        for wn, trans in zip(ir_data["__x__"], ir_data[raw_id], strict=False):
            if pd.isna(wn) or pd.isna(trans):
                continue
            ir_rows.append([wn, trans])
    append_rows(wb.create_sheet("ir_spectrum"), ir_rows)

    hydro_rows = [["component", "content", "unit"]]
    hydro_series = hydro_values.get(raw_id)
    if hydro_series is None:
        hydro_series = hydrocarbon_from_quantified(raw_id)
    if hydro_series is not None:
        for component, content in zip(hydro_components, hydro_series, strict=False):
            if pd.isna(component) or pd.isna(content):
                continue
            hydro_rows.append([str(component).strip(), content, "wt_percent"])
    else:
        for component in hydro_components:
            hydro_rows.append([str(component).strip(), "", "wt_percent"])
    append_rows(wb.create_sheet("hydrocarbon_types"), hydro_rows)

    append_rows(wb.create_sheet("sample_relations"), [["related_sample_id", "relation_type", "description"]])

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    gc_fid_source_id = chromatogram_source_id("gc_fid", raw_id)
    gc_scd_source_id = chromatogram_source_id("gc_scd", raw_id)
    gc_ncd_source_id = chromatogram_source_id("gc_ncd", raw_id)
    return {
        "sample_id": db_id,
        "raw_sample_id": raw_id,
        "file": str(path),
        "has_gc_fid": has_chromatogram(chromatogram_sheet_sets["gc_fid"], gc_fid_source_id),
        "gc_fid_source": gc_fid_source_id if gc_fid_source_id != raw_id else "",
        "has_gc_scd": has_chromatogram(chromatogram_sheet_sets["gc_scd"], gc_scd_source_id),
        "has_gc_ncd": has_chromatogram(chromatogram_sheet_sets["gc_ncd"], gc_ncd_source_id),
        "has_ir": raw_id in ir_data,
        "has_hydrocarbon_types": raw_id in hydro_values or hydrocarbon_from_quantified(raw_id) is not None,
        "hydrocarbon_types_source": "hydrocarbon_types.xlsx" if raw_id in hydro_values else ("quantified_composition" if hydrocarbon_from_quantified(raw_id) is not None else ""),
    }


def main() -> None:
    for path in (GC_FID, GC_SCD, GC_NCD, IR, HYDROCARBON, ELEMENTAL, TEMPLATE):
        if not path.exists():
            raise FileNotFoundError(path)
    ensure_output_root()

    elemental = pd.read_excel(ELEMENTAL)
    elemental["Sample"] = elemental["Sample"].map(canonical_raw_id)
    elemental = elemental.sort_values("Sample", key=lambda col: col.map(sample_sort_key)).reset_index(drop=True)
    raw_ids = elemental["Sample"].astype(str).tolist()

    rng = random.Random(RANDOM_SEED)
    test_count = round(len(raw_ids) * TEST_FRACTION)
    test_samples = sorted(rng.sample(raw_ids, test_count), key=sample_sort_key)
    test_set = set(test_samples)
    train_samples = [sample for sample in raw_ids if sample not in test_set]

    ir_data = load_wide_table(IR)
    hydro_components, hydro_values = load_hydrocarbon_table()
    chromatogram_sheet_sets = {
        "gc_fid": set(pd.ExcelFile(GC_FID).sheet_names),
        "gc_scd": set(pd.ExcelFile(GC_SCD).sheet_names),
        "gc_ncd": set(pd.ExcelFile(GC_NCD).sheet_names),
    }

    manifest: list[dict[str, Any]] = []
    for index, (_, series) in enumerate(elemental.iterrows(), start=1):
        raw_id = canonical_raw_id(series["Sample"])
        split = "test" if raw_id in test_set else "train"
        file_path = OUT_ROOT / split / f"{db_id_for(index)}__{raw_id}.xlsx"
        record = write_sample_workbook(
            path=file_path,
            sample_index=index,
            raw_id=raw_id,
            elemental_row=series,
            ir_data=ir_data,
            hydro_components=hydro_components,
            hydro_values=hydro_values,
            chromatogram_sheet_sets=chromatogram_sheet_sets,
        )
        record["split"] = split
        manifest.append(record)

    split_payload = {
        "schema_version": "petroleum_fraction-experimental-db-split.v1",
        "random_seed": RANDOM_SEED,
        "test_fraction": TEST_FRACTION,
        "sample_count": len(raw_ids),
        "train_count": len(train_samples),
        "test_count": len(test_samples),
        "train_samples": train_samples,
        "test_samples": test_samples,
        "id_policy": {
            "sample_id": "petroleum_fraction_db_### assigned by elemental composition order after raw_sample_id normalization",
            "raw_sample_id": "original petroleum fraction identifier; petroleum_fraction_62_1/2 normalized to petroleum_fraction_62-1/2",
        },
    }
    (OUT_ROOT / "sample_split.json").write_text(json.dumps(split_payload, ensure_ascii=False, indent=2), encoding="utf-8")
    with (OUT_ROOT / "manifest.csv").open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(manifest[0]))
        writer.writeheader()
        writer.writerows(manifest)
    (OUT_ROOT / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"output": str(OUT_ROOT), **{k: split_payload[k] for k in ("sample_count", "train_count", "test_count")}}, ensure_ascii=False))


if __name__ == "__main__":
    main()
